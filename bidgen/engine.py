"""
SMS bid builder — lays scope-of-work pieces onto Google-Sheets-ready workbooks.

One scope = one tab. Every scope tab has the same columns:
  A item | B driver / value | C unit | D factor | E RAW | F QTY | G pkg | H $ each | I line $ | J how it is figured
Blue = type here.  Green = reads another page.  Black = locked formula.  Red = RAW before rounding.

Two ways a scope tab gets built:
  standalone  — the blank scope template.  Anything it needs from the job or from other
                scopes shows as a blue cell on its own page.
  assembled   — inside a full bid.  Job settings read the Job page; numbers that come
                from other scopes read those scopes' tabs (green).

Formula expressions inside scope definitions use placeholders:
  {k}        this tab, key k (value cell)          {k.raw} {k.qty} {k.factor} {k.total} {k.driver}
  {job.k}    a job setting (rates, setup)
  {sum.k}    the total of key k across every OTHER scope on the bid that provides it
"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BLUE = Font(name='Arial', size=10, color='1F4FD8', bold=True)
GREEN = Font(name='Arial', size=10, color='1E7B34')
BLACK = Font(name='Arial', size=10, color='000000')
RED = Font(name='Arial', size=10, color='C00000')
BOLD = Font(name='Arial', size=10, bold=True)
TITLE = Font(name='Arial', size=14, bold=True, color='FFFFFF')
SUB = Font(name='Arial', size=10, italic=True, color='444444')
HEAD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
NOTE = Font(name='Arial', size=9, italic=True, color='333333')
F_TITLE = PatternFill('solid', fgColor='1F2937')
F_HEAD = PatternFill('solid', fgColor='374151')
F_SEC = PatternFill('solid', fgColor='E5E7EB')
F_IN = PatternFill('solid', fgColor='EEF4FF')
F_TOT = PatternFill('solid', fgColor='FFF7D6')
F_BAND = PatternFill('solid', fgColor='DCEBFF')
THIN = Side(style='thin', color='BBBBBB')
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

# Job settings.  key: (label, unit, default, note, group)
JOB = {
    'wage': ('Wage per man / day', '$', 270, 'what a man costs a day', 'rates'),
    'overhead': ('Overhead margin', '', 0.5, 'a margin, so it DIVIDES: 270 ÷ 50% = 540 a man a day', 'rates'),
    'out_of_town': ('OUT OF TOWN?', 'Y / N', 'N', 'Y turns on per diem, hotels and the travel markup', 'rates'),
    'per_diem': ('Per diem $ / man / day', '$', 60, 'only counts when OUT OF TOWN = Y', 'rates'),
    'hotel': ('Hotel $ / room / night', '$', 150, 'two men a room — only counts when OUT OF TOWN = Y', 'rates'),
    'travel_mk': ('Travel markup', '%', 0.3, 'on per diem + hotels', 'rates'),
    'tax': ('Sales tax  (material only)', '%', 0.0825, '8.25% — material only, never labor. Scuppers and collector heads are all-in and never taxed.', 'rates'),
    'mat_markup': ('Material markup  %', '%', None, 'VARIABLE — set it on THIS job. Materials carry no markup until it is answered.', 'rates'),
    'shop_fab': ('Shop fab  $ per sheet', '$', 40, 'every sheet of flat metal gets $40 a sheet fab', 'rates'),
    'field_sf': ('Field  (sq ft)', 'sf', None, 'the net field off the takeoff — every TPO roof page reads it', 'roof'),
    'deck': ('Deck type', 'METAL / WOOD / CONCRETE', None, 'this alone decides the field fastener length and gauge', 'roof'),
    'insul_thk': ('Insulation thickness', 'in', None, '', 'roof'),
    'cover_thk': ('Cover board thickness', 'in', None, '', 'roof'),
    'membrane_thk': ('Membrane thickness', 'mil', None, '', 'roof'),
    'sheet_width': ('Membrane sheet width', 'ft', 5, '5 is the layout that works; 10 is there when it does not. A seam lands every sheet width up the wall — LAW.', 'roof'),
    'iso_ppb': ('ISO plates per board', 'ea', None, 'manufacturer spec', 'roof'),
    'wall_spacing': ('Wall fastener spacing', 'in o.c.', None, 'THE WIND-UPLIFT LEVER — tighten this and only the fastener count moves', 'roof'),
    'bf_spacing': ('Base flashing fastener spacing', 'in o.c.', None, '', 'roof'),
    'field_seam': ('Field seam spacing', 'in o.c.', 12, 'THE WIND-UPLIFT LEVER. 12″ is standard; 6″ doubles every plate in the field.', 'roof'),
}


def money(c):
    c.number_format = '$#,##0.00'


class Row:
    def __init__(self, kind, key=None, **kw):
        self.kind, self.key, self.kw = kind, key, kw
        self.r = None


class Scope:
    """A scope of work: an ordered list of rows plus what it hands to other scopes."""

    def __init__(self, sid, tab, title, subtitle, source, group, provides=None):
        self.sid, self.tab, self.title, self.subtitle = sid, tab, title, subtitle
        self.source, self.group = source, group
        self.rows = []
        self.provides = provides or {}  # name -> placeholder expression of the cell, e.g. '{wall_sq}'

    # ---- row makers -------------------------------------------------
    def h(self, text):
        self.rows.append(Row('h', text=text)); return self

    def note(self, text):
        self.rows.append(Row('note', text=text)); return self

    def cols(self, *heads):
        self.rows.append(Row('cols', heads=heads)); return self

    def inp(self, key, label, unit='', default=None, note=''):
        self.rows.append(Row('inp', key, label=label, unit=unit, default=default, note=note)); return self

    def calc(self, key, label, expr, unit='', note='', fmt=None):
        self.rows.append(Row('calc', key, label=label, expr=expr, unit=unit, note=note, fmt=fmt)); return self

    def link(self, key, label, sumkey, unit='', note=''):
        """Number that comes from other scopes on a full bid; typed here when the page runs alone."""
        self.rows.append(Row('link', key, label=label, sumkey=sumkey, unit=unit, note=note)); return self

    def buy(self, key, label, driver, factor=None, op='x', rnd='up', unit='', pkg='', price=None,
            product='', note='', driver_unit='', taxable=True, raw=None, qty=None, typed_price=None):
        """A line that gets bought.  driver (expr) op factor = RAW → QTY by the rounding rule.
        rnd: 'up' | 'quarter' (fraction <= .25 rounds down) | 'none' | 'walkpad'
        price: key in prices.json (product + $ each) or None."""
        self.rows.append(Row('buy', key, label=label, driver=driver, factor=factor, op=op, rnd=rnd,
                             unit=unit, pkg=pkg, price=price, product=product, note=note,
                             driver_unit=driver_unit, taxable=taxable, raw=raw, qty=qty,
                             typed_price=typed_price)); return self

    def table(self, key, n, cols, total=None, label_fmt='Wall {i}', head=None):
        """Rows 1..n.  cols: list of (colkey, header, kind, expr_or_default) — kind 'in' | 'f' | 'txt'.
        Column order = A,B,C... (first entry is column B; column A is the row label).
        total: dict colkey -> expr for a TOTAL row (expr may use {colkey_range})."""
        self.rows.append(Row('table', key, n=n, cols=cols, total=total or {}, label_fmt=label_fmt, head=head)); return self

    def labor(self, key, label, qty, unit, mode='day', rate=None, note='', travel=True):
        self.rows.append(Row('labor', key, label=label, qty=qty, unit=unit, mode=mode, rate=rate,
                             note=note, travel=travel)); return self

    def check(self, key, label, expr_ok, note=''):
        self.rows.append(Row('check', key, label=label, expr=expr_ok, note=note)); return self

    def blank(self):
        self.rows.append(Row('blank')); return self


# ---------------------------------------------------------------------
class Layout:
    """Places one scope on a worksheet and resolves its placeholders."""

    def __init__(self, scope, ws, mode, ctx):
        self.s, self.ws, self.mode, self.ctx = scope, ws, mode, ctx
        self.cell = {}       # key -> {'v': 'B12', 'raw': 'E12', ...}
        self.local_job = {}  # job key -> cell (standalone)
        self.local_sum = {}  # sum key -> cell (standalone)
        self.buy_rows, self.labor_rows, self.allin_rows = [], [], []

    # placeholder resolution ----------------------------------------
    def ref(self, token):
        if token.startswith('job.'):
            k = token[4:]
            if self.mode == 'assembled':
                return f"Job!$B${self.ctx['job_row'][k]}"
            return self._abs(self.local_job[k])
        if token.startswith('sum.'):
            k = token[4:]
            if self.mode == 'assembled':
                parts = [p for p in self.ctx['providers'].get(k, []) if p[0] != self.s.sid]
                if not parts:
                    return '0'
                return '(' + '+'.join(f"N('{t}'!{c})" for _, t, c in parts) + ')'
            return self._abs(self.local_sum[k])
        if '.' in token:
            k, f = token.split('.', 1)
            return self._abs(self.cell[k][f])
        return self._abs(self.cell[token]['v'])

    @staticmethod
    def _abs(a):
        m = re.match(r'([A-Z]+)(\d+)$', a)
        return f'${m.group(1)}${m.group(2)}'

    def f(self, expr):
        return re.sub(r'\{([a-zA-Z0-9_.]+)\}', lambda m: self.ref(m.group(1)), expr)

    # tokens an expression needs from job / other scopes
    def needs(self):
        job, sums = [], []
        def scan(x):
            if isinstance(x, str):
                for t in re.findall(r'\{([a-zA-Z0-9_.]+)\}', x):
                    if t.startswith('job.') and t[4:] not in job: job.append(t[4:])
                    if t.startswith('sum.') and t[4:] not in sums: sums.append(t[4:])
            elif isinstance(x, (list, tuple)):
                for y in x: scan(y)
            elif isinstance(x, dict):
                for y in x.values(): scan(y)
        for row in self.s.rows:
            scan(row.kw)
        scan(list(self.s.provides.values()))
        # every page carries rates for its own money lines
        for k in ['tax', 'mat_markup']:
            if k not in job: job.append(k)
        if any(r.kind == 'labor' for r in self.s.rows):
            for k in ['wage', 'overhead', 'out_of_town', 'per_diem', 'hotel', 'travel_mk']:
                if k not in job: job.append(k)
        if any(r.kind == 'labor' and r.kw.get('mode') == 'shop' for r in self.s.rows):
            if 'shop_fab' not in job: job.append('shop_fab')
        return job, sums

    # -----------------------------------------------------------------
    def build(self):
        ws, s = self.ws, self.s
        widths = [44, 14, 12, 11, 12, 11, 9, 12, 14, 70]
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.merge_cells('A1:J1'); ws['A1'] = s.title; ws['A1'].font = TITLE; ws['A1'].fill = F_TITLE
        ws.row_dimensions[1].height = 24
        ws.merge_cells('A2:J2'); ws['A2'] = s.subtitle; ws['A2'].font = SUB
        ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 44
        ws.merge_cells('A3:J3')
        ws['A3'] = (f'Came from: {s.source}.   Blue = type here · green = reads another page · black = locked formula · red = RAW before rounding.')
        ws['A3'].font = NOTE

        # summary block rows 5-11 (filled at the end)
        r = 5
        ws.cell(r, 1, 'THIS SCOPE  —  WHAT IT ADDS TO THE BID').font = HEAD
        for c in range(1, 11): ws.cell(r, c).fill = F_HEAD
        self.sum_rows = {}
        for i, (k, lab) in enumerate([('mat', 'Materials  (before tax)'), ('tax', 'Sales tax  (material only)'),
                                       ('mk', 'Material markup'), ('allin', 'All-in items  (never taxed, never marked up)'),
                                       ('labor', 'Labor'), ('travel', 'Travel  (per diem + hotels + travel markup)'),
                                       ('total', 'SCOPE TOTAL')]):
            rr = r + 1 + i
            ws.cell(rr, 1, lab).font = BOLD if k == 'total' else BLACK
            self.sum_rows[k] = rr
        r = r + 9

        job, sums = self.needs()
        if self.mode == 'standalone':
            if job:
                r = self._sec(r, 'RATES & SETUP  —  on a full bid these come from the Job page')
                for k in job:
                    lab, unit, dflt, note, _ = JOB[k]
                    ws.cell(r, 1, lab); c = ws.cell(r, 2, dflt); c.font = BLUE; c.fill = F_IN
                    if unit == '%': c.number_format = '0.00%'
                    ws.cell(r, 3, unit); ws.cell(r, 10, note).font = NOTE
                    self.local_job[k] = f'B{r}'; r += 1
                r += 1
            if sums:
                r = self._sec(r, 'FROM THE OTHER SCOPES  —  on a full bid these fill themselves in')
                for k in sums:
                    lab = self.ctx['sum_labels'].get(k, k)
                    ws.cell(r, 1, lab); c = ws.cell(r, 2); c.font = BLUE; c.fill = F_IN
                    ws.cell(r, 10, 'type it when this page runs alone').font = NOTE
                    self.local_sum[k] = f'B{r}'; r += 1
                r += 1

        # pass 1: assign rows
        rr = r
        for row in s.rows:
            row.r = rr
            if row.kind == 'table':
                rr += 1 + row.kw['n'] + (1 if row.kw['total'] else 0)
                self._table_cells(row)
            elif row.kind == 'labor_head':
                rr += 1
            else:
                self._assign(row)
                rr += 1
        # pass 2: write
        for row in s.rows:
            getattr(self, 'w_' + row.kind)(row)
        end = rr
        self._summary()
        ws.freeze_panes = 'B5'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = 'landscape'
        return end

    # cell bookkeeping
    def _assign(self, row):
        k, r = row.key, row.r
        if not k: return
        if row.kind in ('inp', 'calc', 'link', 'check'):
            self.cell[k] = {'v': f'B{r}'}
        elif row.kind == 'buy':
            self.cell[k] = {'v': f'F{r}', 'driver': f'B{r}', 'factor': f'D{r}', 'raw': f'E{r}',
                            'qty': f'F{r}', 'price': f'H{r}', 'total': f'I{r}'}
        elif row.kind == 'labor':
            self.cell[k] = {'v': f'G{r}', 'men': f'B{r}', 'rate': f'C{r}', 'qty': f'D{r}', 'days': f'F{r}',
                            'labor': f'G{r}', 'pd': f'H{r}', 'hotel': f'I{r}', 'tmk': f'J{r}'}

    def _table_cells(self, row):
        r0 = row.r + 1
        for i in range(1, row.kw['n'] + 1):
            for j, (ck, _, _, _) in enumerate(row.kw['cols']):
                col = get_column_letter(2 + j)
                self.cell[f'{ck}{i}'] = {'v': f'{col}{r0 + i - 1}'}
        for j, (ck, _, _, _) in enumerate(row.kw['cols']):
            col = get_column_letter(2 + j)
            self.cell[f'{ck}_range'] = {'v': f'{col}{r0}'}
            self.cell[f'{ck}_rng'] = None
            self.ctx.setdefault('_ranges', {})[(self.s.sid, ck)] = f'${col}${r0}:${col}${r0 + row.kw["n"] - 1}'
            if ck in row.kw['total']:
                self.cell[f'{ck}_total'] = {'v': f'{col}{r0 + row.kw["n"]}'}

    def _sec(self, r, text):
        self.ws.cell(r, 1, text).font = BOLD
        for c in range(1, 11): self.ws.cell(r, c).fill = F_SEC
        return r + 1

    def _rng(self, expr):
        # {colkey:range} -> absolute range of a table column on this tab
        return re.sub(r'\{([a-zA-Z0-9_]+):range\}',
                      lambda m: self.ctx['_ranges'][(self.s.sid, m.group(1))], expr)

    def F(self, expr):
        return '=' + self.f(self._rng(expr))

    # writers ---------------------------------------------------------
    def w_blank(self, row): pass

    def w_h(self, row):
        self._sec(row.r, row.kw['text'])

    def w_note(self, row):
        ws = self.ws
        ws.merge_cells(start_row=row.r, start_column=1, end_row=row.r, end_column=10)
        t = row.kw['text']
        c = ws.cell(row.r, 1, ('=' + self.f(self._rng(t[1:]))) if t.startswith('=') else t)
        c.font = NOTE; c.fill = F_BAND; c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row.r].height = 30

    def w_cols(self, row):
        for j, h in enumerate(row.kw['heads']):
            c = self.ws.cell(row.r, 1 + j, h); c.font = HEAD; c.fill = F_HEAD

    def w_inp(self, row):
        ws, k = self.ws, row.kw
        ws.cell(row.r, 1, k['label'])
        c = ws.cell(row.r, 2, k['default']); c.font = BLUE; c.fill = F_IN; c.border = BOX
        ws.cell(row.r, 3, k['unit']); ws.cell(row.r, 10, k['note']).font = NOTE

    def w_calc(self, row):
        ws, k = self.ws, row.kw
        ws.cell(row.r, 1, k['label'])
        c = ws.cell(row.r, 2, self.F(k['expr'])); c.font = BLACK
        if k.get('fmt'): c.number_format = k['fmt']
        ws.cell(row.r, 3, k['unit'])
        n = k['note']
        ws.cell(row.r, 10, ('=' + self.f(self._rng(n[1:]))) if n.startswith('=') else n).font = NOTE

    def w_link(self, row):
        ws, k = self.ws, row.kw
        ws.cell(row.r, 1, k['label'])
        if self.mode == 'assembled':
            c = ws.cell(row.r, 2, '=' + self.ref('sum.' + k['sumkey'])); c.font = GREEN
            ws.cell(row.r, 10, k['note']).font = NOTE
        else:
            c = ws.cell(row.r, 2); c.font = BLUE; c.fill = F_IN; c.border = BOX
            ws.cell(row.r, 10, 'on a full bid this fills itself in — ' + k['note']).font = NOTE
        ws.cell(row.r, 3, k['unit'])

    def w_check(self, row):
        ws, k = self.ws, row.kw
        ws.cell(row.r, 1, k['label'])
        c = ws.cell(row.r, 2, '=IFERROR(IF(' + self.f(self._rng(k['expr'])) + ',"OK","CHECK"),"CHECK")')
        c.font = BOLD
        ws.cell(row.r, 10, k['note']).font = NOTE

    def w_table(self, row):
        ws, k = self.ws, row.kw
        r = row.r
        lf = k['label_fmt']
        ws.cell(r, 1, (k.get('head') or (lf.split('{')[0].strip() if isinstance(lf, str) else '') or '')).font = HEAD
        ws.cell(r, 1).fill = F_HEAD
        for j, (_, head, _, _) in enumerate(k['cols']):
            c = ws.cell(r, 2 + j, head); c.font = HEAD; c.fill = F_HEAD
            c.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[r].height = 30
        for i in range(1, k['n'] + 1):
            rr = r + i
            lf = k['label_fmt']
            ws.cell(rr, 1, lf[i - 1] if isinstance(lf, (list, tuple)) else lf.format(i=i))
            for j, (ck, _, kind, ex) in enumerate(k['cols']):
                c = ws.cell(rr, 2 + j)
                if kind == 'in':
                    c.value = ex[i - 1] if isinstance(ex, list) else ex; c.font = BLUE; c.fill = F_IN; c.border = BOX
                elif kind == 'f':
                    c.value = self.F(ex.replace('{i}', str(i)).replace('[i]', str(i))); c.font = BLACK
                else:
                    c.value = ex
        if k['total']:
            rr = r + k['n'] + 1
            ws.cell(rr, 1, 'TOTAL').font = BOLD
            for j, (ck, _, _, _) in enumerate(k['cols']):
                if ck in k['total']:
                    c = ws.cell(rr, 2 + j, self.F(k['total'][ck])); c.font = BOLD; c.fill = F_TOT

    def w_buy(self, row):
        ws, k, r = self.ws, row.kw, row.r
        P = self.ctx['prices'].get(k['price'] or '', {})
        ws.cell(r, 1, k['label'])
        if k['driver']:
            ws.cell(r, 2, self.F(k['driver'])).font = BLACK
        else:
            c = ws.cell(r, 2); c.font = BLUE; c.fill = F_IN; c.border = BOX
        ws.cell(r, 3, k['driver_unit'])
        if k['factor'] is not None:
            c = ws.cell(r, 4, k['factor']); c.font = BLUE; c.fill = F_IN
        B, D = f'B{r}', f'D{r}'
        if k['raw']:
            raw = '=' + self.f(self._rng(k['raw']))
        elif k['op'] == 'x':
            raw = f'=IFERROR(IF(OR({B}="",{D}=""),"",N({B})*N({D})),"")'
        elif k['op'] == '/':
            raw = f'=IFERROR(IF(OR({B}="",{D}="",N({D})=0),"",N({B})/N({D})),"")'
        else:  # '=' driver is the raw
            raw = f'=IFERROR(IF({B}="","",N({B})),"")'
        c = ws.cell(r, 5, raw); c.font = RED; c.number_format = '#,##0.00'
        E = f'E{r}'
        if k['qty']:
            q = '=' + self.f(self._rng(k['qty']))
        elif k['rnd'] == 'quarter':
            q = f'=IFERROR(IF({E}="","",IF(MOD({E},1)<=0.25,ROUNDDOWN({E},0),ROUNDUP({E},0))),"")'
        elif k['rnd'] == 'none':
            q = f'=IFERROR(IF({E}="","",{E}),"")'
        elif k['rnd'] == 'walkpad':
            q = (f'=IFERROR(IF(OR({E}="",N({E})=0),"",INT({E})+IF(MOD({E},1)=0,0,IF(MOD({E},1)<=0.25,0.25,'
                 f'IF(MOD({E},1)<=0.5,0.5,1)))),"")')
        else:
            q = f'=IFERROR(IF({E}="","",ROUNDUP({E},0)),"")'
        ws.cell(r, 6, q).font = BOLD
        ws.cell(r, 7, k['pkg'] or P.get('unit', ''))
        if k['price'] is not None or k['typed_price'] is not None:
            price = k['typed_price'] if k['typed_price'] is not None else P.get('price')
            c = ws.cell(r, 8, price); c.font = BLUE; c.fill = F_IN; money(c)
            c = ws.cell(r, 9, f'=IFERROR(IF(OR(F{r}="",H{r}=""),"",N(F{r})*N(H{r})),"")'); money(c)
            (self.buy_rows if k['taxable'] else self.allin_rows).append(r)
        prod = k['product'] or P.get('product') or ''
        n = k['note']
        if n.startswith('='):
            ws.cell(r, 10, '=' + self.f(self._rng(n[1:])) + (f'&"   ·   {prod}"' if prod else ''))
        else:
            ws.cell(r, 10, (n + ('   ·   ' if n and prod else '') + str(prod)))
        ws.cell(r, 10).font = NOTE

    def w_labor(self, row):
        ws, k, r = self.ws, row.kw, row.r
        J = lambda key: self.ref('job.' + key)
        ws.cell(r, 1, k['label'])
        c = ws.cell(r, 2); c.font = BLUE; c.fill = F_IN; c.border = BOX     # men
        c = ws.cell(r, 3, k['rate']); c.font = BLUE; c.fill = F_IN; c.border = BOX  # per day OR $ per unit
        ws.cell(r, 4, self.F(k['qty'])).font = GREEN
        ws.cell(r, 5, k['unit'])
        crew = f'N(B{r})*IFERROR(N({J("wage")})/N({J("overhead")}),0)'
        if k['mode'] == 'day':
            ws.cell(r, 6, f'=IFERROR(IF(N(C{r})=0,"",N(D{r})/N(C{r})),"")').number_format = '0.00'
            c = ws.cell(r, 7, f'=IFERROR(IF(OR(N(F{r})=0,N(B{r})=0),"",N(F{r})*{crew}),"")'); money(c)
            if k['travel']:
                oot = f'IF(UPPER({J("out_of_town")})="Y",1,0)'
                c = ws.cell(r, 8, f'=IFERROR(IF(OR(N(B{r})=0,N(F{r})=0),"",N(B{r})*N(F{r})*N({J("per_diem")})*{oot}),"")'); money(c)
                c = ws.cell(r, 9, f'=IFERROR(IF(OR(N(B{r})=0,N(F{r})=0),"",ROUNDUP(N(B{r})/2,0)*N(F{r})*N({J("hotel")})*{oot}),"")'); money(c)
                c = ws.cell(r, 10, f'=IFERROR((N(H{r})+N(I{r}))*N({J("travel_mk")}),"")'); money(c)
        elif k['mode'] == 'shop':
            c = ws.cell(r, 3, '=' + J('shop_fab')); c.font = GREEN; money(c)
            c = ws.cell(r, 7, f'=IFERROR(IF(OR(N(D{r})=0,N(C{r})=0),"",N(D{r})*N(C{r})),"")'); money(c)
            ws.cell(r, 2).fill = PatternFill(None)
        else:  # $ per unit
            money(ws.cell(r, 3))
            c = ws.cell(r, 7, f'=IFERROR(IF(OR(N(D{r})=0,N(C{r})=0),"",N(D{r})*N(C{r})),"")'); money(c)
        self.labor_rows.append(r)

    def _summary(self):
        ws, S = self.ws, self.sum_rows
        J = lambda key: self.ref('job.' + key)
        def sumcells(rows, col):
            return '+'.join(f'N({col}{x})' for x in rows) or '0'
        ws.cell(S['mat'], 2, '=' + sumcells(self.buy_rows, 'I'))
        ws.cell(S['tax'], 2, f'=N(B{S["mat"]})*N({J("tax")})')
        ws.cell(S['mk'], 2, f'=(N(B{S["mat"]})+N(B{S["tax"]}))*N({J("mat_markup")})')
        ws.cell(S['allin'], 2, '=' + sumcells(self.allin_rows, 'I'))
        ws.cell(S['labor'], 2, '=' + sumcells(self.labor_rows, 'G'))
        ws.cell(S['travel'], 2, '=' + ('+'.join(f'N(H{x})+N(I{x})+N(J{x})' for x in self.labor_rows) or '0'))
        ws.cell(S['total'], 2, '=' + '+'.join(f'N(B{S[k]})' for k in ['mat', 'tax', 'mk', 'allin', 'labor', 'travel']))
        for k, rr in S.items():
            c = ws.cell(rr, 2); money(c); c.font = BOLD if k == 'total' else BLACK
            if k == 'total':
                for cc in range(1, 3): ws.cell(rr, cc).fill = F_TOT
        ws.cell(S['mat'], 10, 'every priced line on this page added up — the line $ column').font = NOTE
        ws.cell(S['mk'], 10, 'materials w/ tax x material markup — all-in items left out').font = NOTE
        if self.labor_rows:
            hdr = self.labor_rows[0] - 1


def labor_header(scope):
    scope.rows.append(Row('cols', heads=('TRADE', 'MEN', 'PER DAY  /  $ EACH', 'QTY', 'UNIT', 'DAYS', 'LABOR $',
                                         'PER DIEM', 'HOTEL', 'TRAVEL MKUP')))


# ---------------------------------------------------------------------
def build_workbook(scopes, mode, prices, job_values=None, title=None):
    """scopes: list of Scope (order = tab order).  Returns (Workbook, layouts)."""
    wb = Workbook()
    wb.remove(wb.active)
    ctx = {'prices': prices, 'providers': {}, 'job_row': {}, 'sum_labels': {}}
    for s in scopes:
        for name, (lab, _) in s.provides.items():
            ctx['sum_labels'][name] = lab
    # also labels for sums consumed but provided by scopes not in this book
    from .scopes import SUM_LABELS
    for k, v in SUM_LABELS.items(): ctx['sum_labels'].setdefault(k, v)

    if mode == 'assembled':
        wsj = wb.create_sheet('Job')
        _job_page(wsj, ctx, job_values or {}, title)
    layouts = []
    sheets = []
    for s in scopes:
        ws = wb.create_sheet(s.tab)
        lay = Layout(s, ws, mode, ctx)
        layouts.append(lay); sheets.append(ws)
    # pass A: assign cells so providers are known before any formula is written
    for lay in layouts:
        lay._preassign()
    for lay in layouts:
        for name, (_, expr) in lay.s.provides.items():
            ctx['providers'].setdefault(name, []).append((lay.s.sid, lay.s.tab, lay._provided_cell(expr)))
    for lay in layouts:
        lay.build()
    if mode == 'assembled':
        _totals_page(wb, layouts, ctx)
    return wb, layouts


def _job_page(ws, ctx, vals, title):
    ws.column_dimensions['A'].width = 40; ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18; ws.column_dimensions['D'].width = 80
    ws.merge_cells('A1:D1'); ws['A1'] = title or 'JOB'; ws['A1'].font = TITLE; ws['A1'].fill = F_TITLE
    ws['A2'] = 'Typed once here — every scope page reads these. Change one and every page moves.'; ws['A2'].font = SUB
    r = 4
    for grp, head in [('rates', 'RATES'), ('roof', 'THIS ROOF  —  the setup every TPO page reads')]:
        ws.cell(r, 1, head).font = BOLD
        for c in range(1, 5): ws.cell(r, c).fill = F_SEC
        r += 1
        for k, (lab, unit, dflt, note, g) in JOB.items():
            if g != grp: continue
            ws.cell(r, 1, lab)
            c = ws.cell(r, 2, vals.get(k, dflt)); c.font = BLUE; c.fill = F_IN; c.border = BOX
            if unit == '%': c.number_format = '0.00%'
            ws.cell(r, 3, unit); ws.cell(r, 4, note).font = NOTE
            ctx['job_row'][k] = r; r += 1
        r += 1


def _totals_page(wb, layouts, ctx):
    ws = wb.create_sheet('Totals', 1)
    ws.column_dimensions['A'].width = 46; ws.column_dimensions['B'].width = 16
    for col in 'CDEFGH': ws.column_dimensions[col].width = 14
    ws.column_dimensions['I'].width = 60
    ws.merge_cells('A1:I1'); ws['A1'] = 'TOTALS  —  every scope on this bid, one line each'
    ws['A1'].font = TITLE; ws['A1'].fill = F_TITLE
    ws['A2'] = 'Each line reads the THIS SCOPE block at the top of its own page. Nothing on this page is calculated twice.'
    ws['A2'].font = SUB
    heads = ['Scope', 'Materials', 'Sales tax', 'Mat. markup', 'All-in items', 'Labor', 'Travel', 'SCOPE TOTAL', 'Came from']
    for j, h in enumerate(heads):
        c = ws.cell(4, 1 + j, h); c.font = HEAD; c.fill = F_HEAD
    r = 5
    keys = ['mat', 'tax', 'mk', 'allin', 'labor', 'travel', 'total']
    for lay in layouts:
        ws.cell(r, 1, lay.s.title)
        for j, k in enumerate(keys):
            c = ws.cell(r, 2 + j, f"='{lay.s.tab}'!B{lay.sum_rows[k]}"); c.font = GREEN; money(c)
        ws.cell(r, 9, lay.s.source).font = NOTE
        r += 1
    first, last = 5, r - 1
    ws.cell(r, 1, 'SCOPES  —  SUBTOTAL').font = BOLD
    for j in range(7):
        col = get_column_letter(2 + j)
        c = ws.cell(r, 2 + j, f'=SUM({col}{first}:{col}{last})'); c.font = BOLD; money(c); c.fill = F_TOT
    sub = r
    r += 2
    ws.cell(r, 1, 'BEFORE THIS GOES OUT').font = BOLD
    for c in range(1, 10): ws.cell(r, c).fill = F_SEC
    r += 1
    extras = [('Warranty / NDL', None, 'not marked up'), ('Special insurance', None, ''),
              ('Tool purchase / rental', None, ''), ('Bonds  (P&P / performance)', None, ''),
              ('Skytrack rental', None, ''), ('Dumpsters', None, 'tear-off: $700 a dumpster, one load every 10 squares — on the Tear-off page')]
    xr = {}
    for lab, v, note in extras:
        ws.cell(r, 1, lab); c = ws.cell(r, 2, v); c.font = BLUE; c.fill = F_IN; money(c)
        ws.cell(r, 9, note).font = NOTE; xr[lab] = r; r += 1
    # dumpsters come from the tear-off page when it is on the bid
    td = [l for l in layouts if l.s.sid == 'tearoff']
    if td and 'dumpsters' in td[0].cell:
        c = ws.cell(xr['Dumpsters'], 2, f"='{td[0].s.tab}'!{td[0].cell['dumpsters']['v']}"); c.font = GREEN
    ws.cell(r, 1, 'Misc markup  %'); c = ws.cell(r, 2); c.font = BLUE; c.fill = F_IN; c.number_format = '0.00%'
    ws.cell(r, 9, 'VARIABLE — set it on THIS job. Applies to insurance, tools, bonds, Skytrack, dumpsters.').font = NOTE
    mkr = r; r += 1
    ws.cell(r, 1, 'Misc markup')
    c = ws.cell(r, 2, '=(' + '+'.join(f'N(B{xr[k]})' for k in ['Special insurance', 'Tool purchase / rental',
                      'Bonds  (P&P / performance)', 'Skytrack rental', 'Dumpsters']) + f')*N(B{mkr})'); money(c)
    mr = r; r += 2
    ws.cell(r, 1, 'TOTAL BID').font = Font(name='Arial', size=12, bold=True)
    c = ws.cell(r, 2, f'=N(H{sub})+' + '+'.join(f'N(B{v})' for v in xr.values()) + f'+N(B{mr})')
    c.font = Font(name='Arial', size=12, bold=True); money(c); c.fill = F_TOT
    tb = r; r += 1
    ws.cell(r, 1, 'Total squares  (field + kickers)')
    sq = ctx['providers'].get('total_sq', [])
    ws.cell(r, 2, '=' + ('+'.join(f"N('{t}'!{cc})" for _, t, cc in sq) if sq else '0')).font = GREEN
    tq = r; r += 1
    ws.cell(r, 1, 'PRICE PER SQUARE').font = BOLD
    c = ws.cell(r, 2, f'=IFERROR(IF(N(B{tq})=0,"",N(B{tb})/N(B{tq})),"")'); money(c); c.font = BOLD
    r += 2
    # checks
    ws.cell(r, 1, 'CHECKS & BALANCES').font = BOLD
    for c in range(1, 10): ws.cell(r, c).fill = F_SEC
    r += 1
    tabs = {l.s.sid: l for l in layouts}
    if 'walls' in tabs:
        W = tabs['walls']
        terms = [tabs[x] for x in ('coping', 'gravel_stop', 'counter_flashing') if x in tabs]
        parts = []
        for i in range(1, 9):
            p = f"('{W.s.tab}'!{W.cell[f'len{i}']['v']}<>\"\")"
            for t in terms:
                p += f"*('{t.s.tab}'!{t.cell[f'so{i}']['v']}=\"\")"
            parts.append(p)
        ws.cell(r, 1, '1 · Every wall has a termination')
        ws.cell(r, 2, f'=IF(({"+".join(parts)})=0,"OK","CHECK")').font = BOLD
        ws.cell(r, 9, 'a wall with a length but no stretch-out in coping, gravel stop OR counter flashing has no termination assigned').font = NOTE
        r += 1


def _preassign(self):
    """Assign cell addresses for every row (without writing) so other tabs can reference them."""
    # replicate build()'s row placement exactly
    r = 5 + 9
    job, sums = self.needs()
    if self.mode == 'standalone':
        if job: r += 1 + len(job) + 1
        if sums: r += 1 + len(sums) + 1
    self.ctx.setdefault('_ranges', {})
    for row in self.s.rows:
        row.r = r
        if row.kind == 'table':
            self._table_cells(row)
            r += 1 + row.kw['n'] + (1 if row.kw['total'] else 0)
        else:
            self._assign(row); r += 1


def _provided_cell(self, expr):
    m = re.fullmatch(r'\{([a-zA-Z0-9_]+)(?:\.([a-z]+))?\}', expr)
    k, f = m.group(1), m.group(2) or 'v'
    return self._abs(self.cell[k][f])


Layout._preassign = _preassign
Layout._provided_cell = _provided_cell
